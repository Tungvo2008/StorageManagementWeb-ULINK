from __future__ import annotations

import os
import re
import uuid
from datetime import datetime, timezone, date
from decimal import Decimal
from pathlib import Path
from typing import Optional, List
from pydantic import BaseModel

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Request, status
from sqlalchemy import select, func, or_, and_, desc
from sqlalchemy.orm import Session, joinedload

from app.api.deps import get_db, get_current_user, bearer_scheme
from fastapi.security import HTTPAuthorizationCredentials
from app.db.models import (
    CargoBatch, CargoBox, CargoBill, CargoStatement,
    CargoBatchStatus, CargoBoxType, User
)
from app.schemas.trust_cargo import (
    CargoBatchRead, CargoBatchDetail, CargoBatchCreate, CargoBatchUpdate,
    CargoBoxRead, CargoBoxBase, CargoBoxCreate, CargoBoxUpdate,
    CargoBillRead, CargoBillCreate, CargoBillUpdate,
    QuickEntryRequest,
    CargoStatementRead, CargoStatementCreate,
    CaliReconcileRequest, CaliReconcileResponse, CaliReconcileResultItem,
    TrustCargoDashboardSummary, DashboardKpis, DashboardMarketItem, DashboardDirectCustomer,
    DashboardSelfItem, DashboardShipmentLogBatch, DashboardShipmentLogBox
)

router = APIRouter(prefix="/trust-cargo")

BOT_API_KEY = os.environ.get("CARGO_BOT_KEY", "trustcargo_bot_secret_2026")


def get_current_user_or_bot(
    request: Request,
    creds: Optional[HTTPAuthorizationCredentials] = Depends(bearer_scheme),
    db: Session = Depends(get_db),
) -> User:
    token = creds.credentials if creds else request.query_params.get("token")
    if token == BOT_API_KEY:
        bot_user = db.scalar(select(User).where(User.is_active == True)) or db.scalar(select(User))
        if bot_user:
            return bot_user
    if not creds and token:
        creds = HTTPAuthorizationCredentials(scheme="Bearer", credentials=token)
    return get_current_user(request=request, creds=creds, db=db)


def normalize_name(name: str) -> str:
    if not name:
        return ""
    n = name.strip().lower()
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFD", n) if unicodedata.category(c) != "Mn")


def compute_cargo_bill_total(
    billing_type: Optional[str],
    weight_lbs: Optional[Decimal],
    rate_per_lb: Optional[Decimal],
    extra_fee: Optional[Decimal],
    direct_total: Optional[Decimal] = None
) -> tuple[Decimal, Decimal, Decimal]:
    """
    Quy tắc tính cước Trust Cargo:
    - MARKET (Hàng Chợ):
        Đơn giá Chợ = max(0, rate_per_lb - 1)
        Thành tiền = (weight_lbs * Đơn giá Chợ) + extra_fee
    - SELF (Chính mình gửi):
        rate_per_lb = 0, extra_fee = 0, total_amount = 0
    - DIRECT (Khách ngoài):
        Thành tiền = direct_total nếu có > 0, ngược lại = (weight_lbs * rate_per_lb) + extra_fee
    Trả về: (rate_per_lb, extra_fee, total_amount)
    """
    b_type = (billing_type or "MARKET").upper().strip()
    w = weight_lbs or Decimal("0")
    rate = rate_per_lb or Decimal("0")
    extra = extra_fee or Decimal("0")

    if b_type == "SELF":
        return Decimal("0"), Decimal("0"), Decimal("0")
    elif b_type == "MARKET":
        m_rate = max(Decimal("0"), rate - Decimal("1"))
        total = (w * m_rate) + extra
        return rate, extra, total
    else:  # DIRECT
        if direct_total is not None and direct_total > Decimal("0"):
            total = direct_total
        else:
            total = (w * rate) + extra
        return rate, extra, total


# ==================== UPLOAD IMAGE ====================
@router.post("/upload")
async def upload_cargo_image(
    file: UploadFile = File(...),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    upload_dir = Path(__file__).resolve().parents[3] / "assets" / "cargo-images"
    upload_dir.mkdir(parents=True, exist_ok=True)

    ext = Path(file.filename or "image.jpg").suffix.lower()
    if ext not in (".jpg", ".jpeg", ".png", ".webp", ".heic"):
        ext = ".jpg"

    filename = f"{uuid.uuid4().hex[:12]}_{datetime.now().strftime('%Y%m%d%H%M%S')}{ext}"
    file_path = upload_dir / filename

    contents = await file.read()
    with open(file_path, "wb") as f:
        f.write(contents)

    relative_url = f"/api/v1/public/cargo-images/{filename}"
    return {"url": relative_url, "filename": filename}


# ==================== BATCHES ====================
@router.get("/batches", response_model=List[CargoBatchRead])
def list_batches(
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    batches = db.scalars(
        select(CargoBatch).order_by(desc(CargoBatch.ship_date)).offset(skip).limit(limit)
    ).all()

    result = []
    for b in batches:
        b_read = CargoBatchRead.model_validate(b)
        b_read.box_count = len(b.boxes)
        b_read.bill_count = len(b.bills)
        b_read.total_weight_lbs = sum((bill.weight_lbs for bill in b.bills), Decimal("0"))
        b_read.total_amount = sum((bill.total_amount for bill in b.bills), Decimal("0"))
        result.append(b_read)
    return result


@router.post("/batches", response_model=CargoBatchRead)
def create_batch(
    payload: CargoBatchCreate,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    ship_date = payload.ship_date or datetime.now(timezone.utc)
    if not payload.batch_code:
        date_str = ship_date.strftime("%Y%m%d")
        count = db.scalar(select(func.count()).select_from(CargoBatch).where(
            func.date(CargoBatch.ship_date) == ship_date.date()
        )) or 0
        batch_code = f"CARGO-{date_str}-{count + 1:02d}"
    else:
        batch_code = payload.batch_code

    batch = CargoBatch(
        batch_code=batch_code,
        ship_date=ship_date,
        market_name=payload.market_name or "World Market",
        status=CargoBatchStatus.OPEN,
        notes=payload.notes,
    )
    db.add(batch)
    db.commit()
    db.refresh(batch)

    res = CargoBatchRead.model_validate(batch)
    res.box_count = 0
    res.bill_count = 0
    res.total_weight_lbs = Decimal("0")
    res.total_amount = Decimal("0")
    return res


@router.get("/batches/{batch_id}", response_model=CargoBatchDetail)
def get_batch_detail(
    batch_id: int,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    batch = db.scalar(
        select(CargoBatch)
        .options(
            joinedload(CargoBatch.boxes).joinedload(CargoBox.bills),
            joinedload(CargoBatch.bills),
        )
        .where(CargoBatch.id == batch_id)
    )
    if not batch:
        raise HTTPException(status_code=404, detail="Không tìm thấy đợt đóng hàng")

    boxes_read = []
    for box in batch.boxes:
        b_read = CargoBoxRead.model_validate(box)
        b_read.bills = [CargoBillRead.model_validate(b) for b in box.bills]
        boxes_read.append(b_read)

    unassigned_bills = [
        CargoBillRead.model_validate(b) for b in batch.bills if b.box_id is None
    ]

    detail = CargoBatchDetail.model_validate(batch)
    detail.boxes = boxes_read
    detail.unassigned_bills = unassigned_bills
    return detail


@router.put("/batches/{batch_id}", response_model=CargoBatchRead)
def update_batch(
    batch_id: int,
    payload: CargoBatchUpdate,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    batch = db.get(CargoBatch, batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="Không tìm thấy đợt đóng hàng")

    if payload.batch_code is not None:
        batch.batch_code = payload.batch_code
    if payload.ship_date is not None:
        batch.ship_date = payload.ship_date
    if payload.market_name is not None:
        batch.market_name = payload.market_name
    if payload.status is not None:
        batch.status = payload.status
    if payload.notes is not None:
        batch.notes = payload.notes

    db.commit()
    db.refresh(batch)

    res = CargoBatchRead.model_validate(batch)
    res.box_count = len(batch.boxes)
    res.bill_count = len(batch.bills)
    res.total_weight_lbs = sum((bill.weight_lbs for bill in batch.bills), Decimal("0"))
    res.total_amount = sum((bill.total_amount for bill in batch.bills), Decimal("0"))
    return res


@router.delete("/batches/{batch_id}")
def delete_batch(
    batch_id: int,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    batch = db.get(CargoBatch, batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="Không tìm thấy đợt đóng hàng")
    db.delete(batch)
    db.commit()
    return {"status": "ok", "message": f"Đã xoá đợt {batch.batch_code}"}


# ==================== BOXES ====================
@router.post("/batches/{batch_id}/boxes", response_model=CargoBoxRead)
def create_box(
    batch_id: int,
    payload: CargoBoxBase,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    batch = db.get(CargoBatch, batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="Không tìm thấy đợt đóng hàng")

    box = CargoBox(
        batch_id=batch_id,
        box_code=payload.box_code,
        box_type=payload.box_type,
        label_text=payload.label_text,
        box_image_url=payload.box_image_url,
        weight_lbs=payload.weight_lbs,
        notes=payload.notes,
        is_hub=payload.is_hub,
        hub_consignee_name=payload.hub_consignee_name,
        tracking_number=payload.tracking_number,
        delivery_status=payload.delivery_status or "PENDING",
        delivered_at=payload.delivered_at,
        cali_code=payload.cali_code,
    )
    db.add(box)
    db.commit()
    db.refresh(box)
    return CargoBoxRead.model_validate(box)


@router.put("/boxes/{box_id}", response_model=CargoBoxRead)
def update_box(
    box_id: int,
    payload: CargoBoxUpdate,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    box = db.get(CargoBox, box_id)
    if not box:
        raise HTTPException(status_code=404, detail="Không tìm thấy thùng")

    if payload.box_code is not None:
        box.box_code = payload.box_code
    if payload.box_type is not None:
        box.box_type = payload.box_type
    if payload.label_text is not None:
        box.label_text = payload.label_text
    if payload.box_image_url is not None:
        box.box_image_url = payload.box_image_url
    if payload.weight_lbs is not None:
        box.weight_lbs = payload.weight_lbs
    if payload.notes is not None:
        box.notes = payload.notes
    if payload.is_hub is not None:
        box.is_hub = payload.is_hub
    if payload.hub_consignee_name is not None:
        box.hub_consignee_name = payload.hub_consignee_name
    if payload.tracking_number is not None:
        box.tracking_number = payload.tracking_number
    if payload.delivery_status is not None:
        box.delivery_status = payload.delivery_status
    if payload.delivered_at is not None:
        box.delivered_at = payload.delivered_at
    if payload.cali_code is not None:
        box.cali_code = payload.cali_code

    db.commit()
    db.refresh(box)
    b_read = CargoBoxRead.model_validate(box)
    b_read.bills = [CargoBillRead.model_validate(b) for b in box.bills]
    return b_read


@router.delete("/boxes/{box_id}")
def delete_box(
    box_id: int,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    box = db.get(CargoBox, box_id)
    if not box:
        raise HTTPException(status_code=404, detail="Không tìm thấy thùng")
    for b in box.bills:
        b.box_id = None
    db.delete(box)
    db.commit()
    return {"status": "ok", "message": "Đã xoá thùng thành công"}


# ==================== BILLS ====================
@router.post("/bills", response_model=CargoBillRead)
def create_bill(
    payload: CargoBillCreate,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    batch = db.get(CargoBatch, payload.batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="Không tìm thấy đợt đóng hàng")

    b_type = (payload.billing_type or "MARKET").upper().strip()
    rate, extra, tot = compute_cargo_bill_total(
        b_type, payload.weight_lbs, payload.rate_per_lb, payload.extra_fee, payload.total_amount if b_type == "DIRECT" else None
    )
    bill = CargoBill(
        batch_id=payload.batch_id,
        box_id=payload.box_id,
        invoice_no=payload.invoice_no,
        bill_date=payload.bill_date or datetime.now(timezone.utc),
        shipper_name=payload.shipper_name,
        shipper_phone=payload.shipper_phone,
        shipper_address=payload.shipper_address,
        consignee_name=payload.consignee_name,
        consignee_phone=payload.consignee_phone,
        consignee_address=payload.consignee_address,
        weight_lbs=payload.weight_lbs,
        declared_value=payload.declared_value,
        commodities=payload.commodities,
        rate_per_lb=rate,
        extra_fee=extra,
        total_amount=tot,
        bill_image_url=payload.bill_image_url,
        market_name=payload.market_name or batch.market_name,
        is_paid=payload.is_paid,
        billing_type=b_type,
        packed_date=payload.packed_date or datetime.now(timezone.utc),
        cali_tracking_code=payload.cali_tracking_code,
        cali_weight_lbs=payload.cali_weight_lbs,
        cali_amount=payload.cali_amount,
        reconcile_status=payload.reconcile_status,
        notes=payload.notes,
    )
    db.add(bill)
    db.commit()
    db.refresh(bill)
    return CargoBillRead.model_validate(bill)


class BillPaidStatusUpdate(BaseModel):
    is_paid: bool


@router.patch("/bills/{bill_id}/paid")
def update_bill_paid_status(
    bill_id: int,
    payload: BillPaidStatusUpdate,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    bill = db.get(CargoBill, bill_id)
    if not bill:
        raise HTTPException(status_code=404, detail="Không tìm thấy bill")
    bill.is_paid = payload.is_paid
    db.commit()
    db.refresh(bill)
    return {"status": "ok", "bill_id": bill.id, "is_paid": bill.is_paid}


@router.put("/bills/{bill_id}", response_model=CargoBillRead)
def update_bill(
    bill_id: int,
    payload: CargoBillUpdate,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    bill = db.get(CargoBill, bill_id)
    if not bill:
        raise HTTPException(status_code=404, detail="Không tìm thấy bill")

    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(bill, k, v)

    rate, extra, tot = compute_cargo_bill_total(
        bill.billing_type, bill.weight_lbs, bill.rate_per_lb, bill.extra_fee, bill.total_amount if (bill.billing_type or "").upper() == "DIRECT" else None
    )
    bill.rate_per_lb = rate
    bill.extra_fee = extra
    bill.total_amount = tot

    db.commit()
    db.refresh(bill)
    return CargoBillRead.model_validate(bill)


@router.delete("/bills/{bill_id}")
def delete_bill(
    bill_id: int,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    bill = db.get(CargoBill, bill_id)
    if not bill:
        raise HTTPException(status_code=404, detail="Không tìm thấy bill")
    db.delete(bill)
    db.commit()
    return {"status": "ok", "message": "Đã xoá bill thành công"}


# ==================== QUICK ENTRY FOR BOT ====================
@router.post("/quick-entry")
def quick_entry(
    payload: QuickEntryRequest,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    today_dt = datetime.now(timezone.utc)
    target_date = today_dt.date()
    if payload.batch_date:
        try:
            target_date = datetime.strptime(payload.batch_date, "%Y-%m-%d").date()
        except ValueError:
            pass

    market = payload.market_name or "World Market"
    batch = db.scalar(
        select(CargoBatch)
        .where(
            and_(
                func.date(CargoBatch.ship_date) == target_date,
                CargoBatch.market_name == market,
            )
        )
        .order_by(desc(CargoBatch.id))
    )

    if not batch:
        date_str = target_date.strftime("%Y%m%d")
        batch = CargoBatch(
            batch_code=f"CARGO-{date_str}-{market[:4].upper()}",
            ship_date=datetime(target_date.year, target_date.month, target_date.day, 12, 0, 0, tzinfo=timezone.utc),
            market_name=market,
            status=CargoBatchStatus.OPEN,
            notes="Tự động tạo bởi Zen Bot",
        )
        db.add(batch)
        db.commit()
        db.refresh(batch)

    if payload.entry_type == "BILL":
        b_data = payload.bill
        if not b_data or not b_data.consignee_name:
            raise HTTPException(status_code=400, detail="Thiếu thông tin người nhận (consignee_name)")

        matched_box_id = b_data.box_id
        if not matched_box_id:
            norm_consignee = normalize_name(b_data.consignee_name)
            consignee_words = [p for p in norm_consignee.split() if len(p) >= 2]

            # 1. Exact full name match in label
            for box in batch.boxes:
                norm_label = normalize_name(box.label_text or "")
                if norm_consignee and norm_consignee in norm_label:
                    matched_box_id = box.id
                    break

            # 2. Phone match if available
            if not matched_box_id and b_data.consignee_phone:
                clean_bill_phone = re.sub(r"\D", "", b_data.consignee_phone)
                if len(clean_bill_phone) >= 7:
                    for box in batch.boxes:
                        clean_box_label = re.sub(r"\D", "", box.label_text or "")
                        if clean_bill_phone[-7:] in clean_box_label:
                            matched_box_id = box.id
                            break

            # 3. Best multi-word match (at least 2 words match, or all words if < 2)
            if not matched_box_id and consignee_words:
                best_score = 0
                best_box_id = None
                for box in batch.boxes:
                    norm_label = normalize_name(box.label_text or "")
                    box_words = set(norm_label.split())
                    score = sum(1 for w in consignee_words if w in box_words)
                    min_required = min(2, len(consignee_words))
                    if score >= min_required and score > best_score:
                        best_score = score
                        best_box_id = box.id
                if best_box_id:
                    matched_box_id = best_box_id

        # Auto-match into open Hub Box (Thao Tran) if not matched to any specific box
        if not matched_box_id:
            hub_boxes = [box for box in batch.boxes if box.is_hub]
            if len(hub_boxes) == 1:
                matched_box_id = hub_boxes[0].id
                b_data.is_hub_item = True
                if not b_data.hub_consignee_name:
                    b_data.hub_consignee_name = hub_boxes[0].hub_consignee_name or "Thảo Trần"

        is_hub_item = bool(b_data.is_hub_item)
        hub_consignee = b_data.hub_consignee_name

        if matched_box_id:
            box = db.get(CargoBox, matched_box_id)
            if box:
                if len(box.bills) + 1 >= 2 or box.is_hub:
                    box.box_type = CargoBoxType.GROUPED
                if box.is_hub:
                    is_hub_item = True
                    if not hub_consignee:
                        hub_consignee = box.hub_consignee_name or (box.label_text.splitlines()[0] if box.label_text else None)

        b_type = (b_data.billing_type or "MARKET").upper().strip()
        rate, extra, tot = compute_cargo_bill_total(
            b_type, b_data.weight_lbs, b_data.rate_per_lb, b_data.extra_fee, b_data.total_amount if b_type == "DIRECT" else None
        )
        bill = CargoBill(
            batch_id=batch.id,
            box_id=matched_box_id,
            invoice_no=b_data.invoice_no or "",
            bill_date=b_data.bill_date or datetime.now(timezone.utc),
            packed_date=b_data.packed_date or datetime.now(timezone.utc),
            billing_type=b_type,
            shipper_name=b_data.shipper_name or "",
            shipper_phone=b_data.shipper_phone or "",
            shipper_address=b_data.shipper_address or "",
            consignee_name=b_data.consignee_name,
            consignee_phone=b_data.consignee_phone or "",
            consignee_address=b_data.consignee_address or "",
            weight_lbs=b_data.weight_lbs,
            declared_value=b_data.declared_value,
            commodities=b_data.commodities or "",
            rate_per_lb=rate,
            extra_fee=extra,
            total_amount=tot,
            bill_image_url=b_data.bill_image_url,
            market_name=market,
            notes=b_data.notes,
            is_hub_item=is_hub_item,
            hub_consignee_name=hub_consignee,
        )
        db.add(bill)
        db.commit()
        db.refresh(bill)

        return {
            "status": "ok",
            "type": "BILL",
            "bill_id": bill.id,
            "batch_id": batch.id,
            "matched_box_id": matched_box_id,
            "consignee_name": bill.consignee_name,
            "weight_lbs": float(bill.weight_lbs),
            "total_amount": float(bill.total_amount),
        }

    elif payload.entry_type == "BOX":
        bx_data = payload.box
        if not bx_data:
            raise HTTPException(status_code=400, detail="Thiếu thông tin thùng hàng (box)")

        # Auto-detect Thao Tran Hub box from label or phone
        lbl_norm = normalize_name(bx_data.label_text or "")
        clean_lbl_phone = re.sub(r"\D", "", bx_data.label_text or "")
        is_thao_tran_hub = "thao tran" in lbl_norm or "0933360174" in clean_lbl_phone

        is_hub_val = bool(bx_data.is_hub or bx_data.hub_consignee_name or is_thao_tran_hub)
        hub_name_val = bx_data.hub_consignee_name or ("Thảo Trần" if is_thao_tran_hub else (bx_data.label_text.strip() if is_hub_val else None))
        if is_hub_val and (not hub_name_val or not hub_name_val.strip()):
            hub_name_val = "Thảo Trần"

        box_code = bx_data.box_code
        if not box_code:
            existing_count = len(batch.boxes)
            box_code = f"THUNG-{existing_count + 1:02d}"

        box = CargoBox(
            batch_id=batch.id,
            box_code=box_code,
            box_type=CargoBoxType.GROUPED if is_hub_val else bx_data.box_type,
            label_text=bx_data.label_text,
            box_image_url=bx_data.box_image_url,
            weight_lbs=bx_data.weight_lbs,
            is_hub=is_hub_val,
            hub_consignee_name=hub_name_val,
            notes=bx_data.notes,
        )
        db.add(box)
        db.commit()
        db.refresh(box)

        created_lines = []
        if bx_data.lines:
            for line in bx_data.lines:
                if not line.consignee_name or not line.consignee_name.strip():
                    raise HTTPException(status_code=422, detail="Mỗi kiện/line trong thùng bắt buộc có Tên người nhận (consignee_name)")
                if not line.consignee_phone or not line.consignee_phone.strip():
                    raise HTTPException(status_code=422, detail=f"Khách '{line.consignee_name}' bắt buộc có Số điện thoại người nhận (consignee_phone)")
                if not line.consignee_address or not line.consignee_address.strip():
                    raise HTTPException(status_code=422, detail=f"Khách '{line.consignee_name}' bắt buộc có Địa chỉ người nhận tại VN (consignee_address)")
                if line.weight_lbs is None or line.weight_lbs <= Decimal("0"):
                    raise HTTPException(status_code=422, detail=f"Khách '{line.consignee_name}' bắt buộc có cân nặng (weight_lbs > 0)")

                b_type = (line.billing_type or "MARKET").upper().strip()
                if b_type == "DIRECT":
                    if not line.shipper_name or not line.shipper_name.strip():
                        raise HTTPException(status_code=422, detail=f"Khách ngoài '{line.consignee_name}' bắt buộc có Tên người gửi để xuất bill thu tiền")

                    if (line.total_amount or 0) <= 0 and (line.rate_per_lb or 0) <= 0:
                        raise HTTPException(status_code=422, detail=f"Khách ngoài '{line.consignee_name}' bắt buộc có đơn giá hoặc tổng cước > 0")
                elif b_type == "SELF":
                    line.total_amount = Decimal("0")
                    line.rate_per_lb = Decimal("0")
                    line.extra_fee = Decimal("0")
                elif b_type == "MARKET":
                    if not line.bill_image_url and not line.invoice_no:
                        raise HTTPException(status_code=422, detail=f"Hàng chợ '{line.consignee_name}' bắt buộc có ảnh chụp bill chợ hoặc số hóa đơn (invoice_no)")

                rate, extra, tot = compute_cargo_bill_total(
                    b_type, line.weight_lbs, line.rate_per_lb, line.extra_fee, line.total_amount if b_type == "DIRECT" else None
                )
                bill_item = CargoBill(
                    batch_id=batch.id,
                    box_id=box.id,
                    invoice_no=line.invoice_no or "",
                    bill_date=datetime.now(timezone.utc),
                    packed_date=datetime.now(timezone.utc),
                    billing_type=b_type,
                    shipper_name=line.shipper_name or "",
                    shipper_phone=line.shipper_phone or "",
                    consignee_name=line.consignee_name.strip(),
                    consignee_phone=line.consignee_phone.strip(),
                    consignee_address=line.consignee_address.strip(),
                    weight_lbs=line.weight_lbs,
                    rate_per_lb=rate,
                    extra_fee=extra,
                    total_amount=tot,
                    commodities=line.commodities or "",
                    bill_image_url=line.bill_image_url or bx_data.box_image_url,
                    market_name=market,
                    hub_consignee_name=box.hub_consignee_name,
                    is_hub_item=box.is_hub,
                    notes=line.notes,
                )
                db.add(bill_item)
                created_lines.append({
                    "consignee_name": line.consignee_name,
                    "weight_lbs": float(line.weight_lbs),
                    "billing_type": b_type
                })

            if len(bx_data.lines) >= 2:
                box.box_type = CargoBoxType.GROUPED

        matched_bills_count = len(created_lines)
        if not created_lines:
            norm_label = normalize_name(box.label_text)
            unassigned = [b for b in batch.bills if b.box_id is None]
            for b in unassigned:
                norm_b = normalize_name(b.consignee_name)
                parts = [p for p in norm_b.split() if len(p) >= 2]
                is_match = norm_b in norm_label or any(p in norm_label for p in parts)
                if is_match:
                    b.box_id = box.id
                    matched_bills_count += 1
            if matched_bills_count >= 2:
                box.box_type = CargoBoxType.GROUPED

        db.commit()
        return {
            "status": "ok",
            "type": "BOX",
            "box_id": box.id,
            "box_code": box.box_code,
            "batch_id": batch.id,
            "box_type": box.box_type.value,
            "created_lines": created_lines,
            "matched_bills_count": matched_bills_count,
        }

    raise HTTPException(status_code=400, detail="entry_type phải là 'BILL' hoặc 'BOX'")


# ==================== MARKET STATEMENT CALCULATION & SAVED BILLS ====================
@router.get("/reports/market-statement")
def get_market_statement(
    market_name: Optional[str] = None,
    filter_mode: str = "UNBILLED",  # "UNBILLED" (từ đợt gần nhất đến nay) or "ALL" (toàn bộ)
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    """
    Quy tắc tính tiền Chợ:
    - Đơn giá cho Chợ = max(0, rate_per_lb - 1)
    - Lấy hết 100% Phụ thu
    - Thành tiền = (Cân nặng * Đơn giá Chợ) + Phụ thu
    """
    query = select(CargoBill).where(CargoBill.billing_type == "MARKET").order_by(desc(CargoBill.bill_date))
    if market_name:
        query = query.where(CargoBill.market_name.ilike(f"%{market_name}%"))

    if filter_mode == "UNBILLED":
        # Chỉ lấy các bill chưa gán vào statement nào
        query = query.where(CargoBill.statement_id == None)

    if start_date:
        try:
            s_dt = datetime.strptime(start_date, "%Y-%m-%d")
            query = query.where(CargoBill.bill_date >= s_dt)
        except ValueError:
            pass
    if end_date:
        try:
            e_dt = datetime.strptime(end_date, "%Y-%m-%d")
            query = query.where(CargoBill.bill_date <= e_dt)
        except ValueError:
            pass

    bills = db.scalars(query).all()

    items = []
    total_weight = Decimal("0")
    total_extra = Decimal("0")
    total_market_amount = Decimal("0")

    for idx, b in enumerate(bills):
        w = b.weight_lbs or Decimal("0")
        # Công thức: Đơn giá chợ = đơn giá gốc - 1
        raw_rate = b.rate_per_lb or Decimal("0")
        market_rate = max(Decimal("0"), raw_rate - Decimal("1"))
        extra = b.extra_fee or Decimal("0")
        line_total = (w * market_rate) + extra

        total_weight += w
        total_extra += extra
        total_market_amount += line_total

        items.append({
            "stt": idx + 1,
            "id": b.id,
            "date": f"{b.bill_date.month}/{b.bill_date.day}/{b.bill_date.year}",
            "consignee_name": b.consignee_name,
            "weight_lbs": float(w),
            "rate_per_lb": float(market_rate),  # Đơn giá đã trừ 1
            "original_rate": float(raw_rate),
            "extra_fee": float(extra),
            "total_amount": float(line_total),
            "market_name": b.market_name,
            "bill_image_url": b.bill_image_url,
            "is_billed": b.statement_id is not None,
            "statement_id": b.statement_id,
        })

    return {
        "market_name": market_name or "Tất cả các Chợ",
        "filter_mode": filter_mode,
        "date_title": datetime.now(timezone.utc).strftime("%-m/%-d/%Y"),
        "total_bills": len(items),
        "total_weight_lbs": float(total_weight),
        "total_extra_fee": float(total_extra),
        "total_amount": float(total_market_amount),
        "bills": items,
    }



def generate_market_statement_excel(
    statement_title: str,
    date_title: str,
    market_name: str,
    bills: list[dict],
    total_amount: float,
    total_weight: float,
    total_extra: float,
) -> io.BytesIO:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bill Vận Chuyển"
    ws.views.sheetView[0].showGridLines = True

    thin_side = Side(border_style="thin", color="000000")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    # Title & Info
    ws["A1"] = statement_title
    ws["A1"].font = Font(name="Arial", size=16, bold=True)

    ws["A2"] = f"Date: {date_title}"
    ws["A2"].font = Font(name="Arial", size=11, bold=True)



    # Table Header at Row 5
    headers = ["Stt", "Tên Người Nhận", "Ngày gửi hàng", "Cân nặng (lbs)", "Phí ($)", "Phụ Thu ($)", "Tổng Tiền ($)"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_idx, value=h)
        cell.font = Font(name="Arial", size=11, bold=True)
        cell.fill = header_fill
        cell.border = thin_border
        if col_idx in [1, 3]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in [4, 5, 6, 7]:
            cell.alignment = Alignment(horizontal="right", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

    current_row = 6
    for b in bills:
        row_vals = [
            b.get("stt"),
            b.get("consignee_name") or "",
            b.get("date") or "",
            float(b.get("weight_lbs") or 0),
            float(b.get("rate_per_lb") or 0),
            float(b.get("extra_fee") or 0),
            float(b.get("total_amount") or 0),
        ]
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = Font(name="Arial", size=10)
            cell.border = thin_border
            if col_idx in [1, 3]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [4, 5, 6, 7]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if col_idx in [5, 6, 7]:
                    cell.number_format = "$#,##0.00"
                elif col_idx == 4:
                    cell.number_format = "#,##0.0"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1

    # Total Row
    cell_total_label = ws.cell(row=current_row, column=1, value="Total:")
    cell_total_label.font = Font(name="Arial", size=11, bold=True)
    cell_total_label.fill = total_fill
    cell_total_label.border = thin_border

    for c in range(2, 4):
        blank_cell = ws.cell(row=current_row, column=c, value="")
        blank_cell.fill = total_fill
        blank_cell.border = thin_border

    # Weight total
    cell_w = ws.cell(row=current_row, column=4, value=float(total_weight))
    cell_w.font = Font(name="Arial", size=11, bold=True)
    cell_w.fill = total_fill
    cell_w.alignment = Alignment(horizontal="right", vertical="center")
    cell_w.border = thin_border
    cell_w.number_format = "#,##0.0"

    # Rate blank
    cell_rate = ws.cell(row=current_row, column=5, value="")
    cell_rate.fill = total_fill
    cell_rate.border = thin_border

    # Extra fee total
    cell_extra = ws.cell(row=current_row, column=6, value=float(total_extra))
    cell_extra.font = Font(name="Arial", size=11, bold=True)
    cell_extra.fill = total_fill
    cell_extra.alignment = Alignment(horizontal="right", vertical="center")
    cell_extra.border = thin_border
    cell_extra.number_format = "$#,##0.00"

    # Total amount
    cell_tot = ws.cell(row=current_row, column=7, value=float(total_amount))
    cell_tot.font = Font(name="Arial", size=12, bold=True, color="0F766E")
    cell_tot.fill = total_fill
    cell_tot.alignment = Alignment(horizontal="right", vertical="center")
    cell_tot.border = thin_border
    cell_tot.number_format = "$#,##0.00"

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.get("/reports/market-statement/export")
def export_market_statement_excel(
    market_name: Optional[str] = None,
    filter_mode: str = "UNBILLED",
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    stmt = get_market_statement(
        market_name=market_name,
        filter_mode=filter_mode,
        start_date=start_date,
        end_date=end_date,
        db=db,
        user=user,
    )
    date_clean = datetime.now(timezone.utc).strftime("%Y%m%d")
    buf = generate_market_statement_excel(
        statement_title="Bill Vận Chuyển",
        date_title=stmt["date_title"],
        market_name=stmt["market_name"],
        bills=stmt["bills"],
        total_amount=stmt["total_amount"],
        total_weight=stmt["total_weight_lbs"],
        total_extra=stmt["total_extra_fee"],
    )
    safe_market = (stmt["market_name"] or "Market").replace(" ", "_")
    filename = f"Bill_Van_Chuyen_{safe_market}_{date_clean}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/statements/{statement_id}/export")
def export_saved_statement_excel(
    statement_id: int,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    statement = db.scalar(
        select(CargoStatement)
        .options(joinedload(CargoStatement.bills))
        .where(CargoStatement.id == statement_id)
    )
    if not statement:
        raise HTTPException(status_code=404, detail="Không tìm thấy bảng kê")

    bills_data = []
    for idx, b in enumerate(statement.bills or []):
        w = b.weight_lbs or Decimal("0")
        raw_rate = b.rate_per_lb or Decimal("0")
        m_rate = max(Decimal("0"), raw_rate - Decimal("1"))
        extra = b.extra_fee or Decimal("0")
        tot = (w * m_rate) + extra
        bills_data.append({
            "stt": idx + 1,
            "consignee_name": b.consignee_name,
            "date": f"{b.bill_date.month}/{b.bill_date.day}/{b.bill_date.year}" if b.bill_date else "",
            "weight_lbs": float(w),
            "rate_per_lb": float(m_rate),
            "extra_fee": float(extra),
            "total_amount": float(tot),
        })

    date_str = f"{statement.statement_date.month}/{statement.statement_date.day}/{statement.statement_date.year}" if statement.statement_date else ""
    buf = generate_market_statement_excel(
        statement_title=f"Bill Vận Chuyển — {statement.statement_code}",
        date_title=date_str,
        market_name=statement.market_name or "World Market",
        bills=bills_data,
        total_amount=float(statement.total_amount or 0),
        total_weight=float(statement.total_weight_lbs or 0),
        total_extra=float(statement.total_extra_fee or 0),
    )
    filename = f"Bill_{statement.statement_code}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )



def generate_market_statement_pdf(
    statement_title: str,
    date_title: str,
    market_name: str,
    bills: list[dict],
    total_amount: float,
    total_weight: float,
    total_extra: float,
) -> io.BytesIO:
    import html
    import io
    from pathlib import Path
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, portrait
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    fonts = Path(__file__).resolve().parents[3] / "assets" / "fonts"
    pdfmetrics.registerFont(TTFont("DejaVu", str(fonts / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVuBold", str(fonts / "DejaVuSans-Bold.ttf")))

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="TitleStyle", fontName="DejaVuBold", fontSize=18, leading=22, textColor=colors.HexColor("#000000")))
    styles.add(ParagraphStyle(name="SubTitleStyle", fontName="DejaVuBold", fontSize=12, leading=16, alignment=2))
    styles.add(ParagraphStyle(name="MarketStyle", fontName="DejaVu", fontSize=10, leading=14, textColor=colors.HexColor("#475569")))

    styles.add(ParagraphStyle(name="TH_Center", fontName="DejaVuBold", fontSize=8.5, leading=11, alignment=1))
    styles.add(ParagraphStyle(name="TH_Left", fontName="DejaVuBold", fontSize=8.5, leading=11, alignment=0))
    styles.add(ParagraphStyle(name="TH_Right", fontName="DejaVuBold", fontSize=8.5, leading=11, alignment=2))

    styles.add(ParagraphStyle(name="TD_Center", fontName="DejaVu", fontSize=8, leading=10, alignment=1))
    styles.add(ParagraphStyle(name="TD_Left", fontName="DejaVu", fontSize=8, leading=10, alignment=0))
    styles.add(ParagraphStyle(name="TD_Right", fontName="DejaVu", fontSize=8, leading=10, alignment=2))

    styles.add(ParagraphStyle(name="TF_Label", fontName="DejaVuBold", fontSize=9, leading=11, alignment=0))
    styles.add(ParagraphStyle(name="TF_Right", fontName="DejaVuBold", fontSize=9, leading=11, alignment=2))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=portrait(letter),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24
    )

    story = []

    header_table = Table(
        [
            [Paragraph(html.escape(statement_title), styles["TitleStyle"]), Paragraph(f"Date: {html.escape(date_title)}", styles["SubTitleStyle"])]
        ],
        colWidths=[360, 204]
    )
    header_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'BOTTOM'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ('TOPPADDING', (0,0), (-1,-1), 0),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 4))
    story.append(Spacer(1, 12))

    col_widths = [28, 180, 76, 70, 55, 75, 80]

    table_data = [
        [
            Paragraph("Stt", styles["TH_Center"]),
            Paragraph("Tên Người Nhận", styles["TH_Left"]),
            Paragraph("Ngày gửi hàng", styles["TH_Center"]),
            Paragraph("Cân nặng (lbs)", styles["TH_Right"]),
            Paragraph("Phí ($)", styles["TH_Right"]),
            Paragraph("Phụ Thu ($)", styles["TH_Right"]),
            Paragraph("Tổng Tiền ($)", styles["TH_Right"]),
        ]
    ]

    for b in bills:
        stt = str(b.get("stt") or "")
        name = str(b.get("consignee_name") or "")
        dt = str(b.get("date") or "")
        w = f"{float(b.get('weight_lbs') or 0):.1f}"
        rate = f"{float(b.get('rate_per_lb') or 0):.2f}"
        extra = f"{float(b.get('extra_fee') or 0):.2f}"
        tot = f"{float(b.get('total_amount') or 0):.2f}"

        table_data.append([
            Paragraph(html.escape(stt), styles["TD_Center"]),
            Paragraph(html.escape(name), styles["TD_Left"]),
            Paragraph(html.escape(dt), styles["TD_Center"]),
            Paragraph(w, styles["TD_Right"]),
            Paragraph(rate, styles["TD_Right"]),
            Paragraph(extra, styles["TD_Right"]),
            Paragraph(tot, styles["TD_Right"]),
        ])

    tot_w_str = f"{float(total_weight or 0):.1f}"
    tot_extra_str = f"{float(total_extra or 0):.2f}"
    tot_amt_str = f"{float(total_amount or 0):.2f}"

    table_data.append([
        Paragraph("Total:", styles["TF_Label"]),
        Paragraph("", styles["TD_Left"]),
        Paragraph("", styles["TD_Center"]),
        Paragraph(tot_w_str, styles["TF_Right"]),
        Paragraph("", styles["TD_Right"]),
        Paragraph(tot_extra_str, styles["TF_Right"]),
        Paragraph(f"${tot_amt_str}", styles["TF_Right"]),
    ])

    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BOX', (0,0), (-1,-1), 1, colors.black),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.black),
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#F8FAFC")),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor("#F1F5F9")),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
    ]))

    story.append(t)
    doc.build(story)
    buf.seek(0)
    return buf


@router.get("/reports/market-statement/export-pdf")
def export_market_statement_pdf(
    market_name: Optional[str] = None,
    filter_mode: str = "UNBILLED",
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    stmt = get_market_statement(
        market_name=market_name,
        filter_mode=filter_mode,
        start_date=start_date,
        end_date=end_date,
        db=db,
        user=user,
    )
    date_clean = datetime.now(timezone.utc).strftime("%Y%m%d")
    buf = generate_market_statement_pdf(
        statement_title="Bill Vận Chuyển",
        date_title=stmt["date_title"],
        market_name=stmt["market_name"],
        bills=stmt["bills"],
        total_amount=stmt["total_amount"],
        total_weight=stmt["total_weight_lbs"],
        total_extra=stmt["total_extra_fee"],
    )
    safe_market = (stmt["market_name"] or "Market").replace(" ", "_")
    filename = f"Bill_Van_Chuyen_{safe_market}_{date_clean}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/statements/{statement_id}/export-pdf")
def export_saved_statement_pdf(
    statement_id: int,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    statement = db.scalar(
        select(CargoStatement)
        .options(joinedload(CargoStatement.bills))
        .where(CargoStatement.id == statement_id)
    )
    if not statement:
        raise HTTPException(status_code=404, detail="Không tìm thấy bảng kê")

    bills_data = []
    for idx, b in enumerate(statement.bills or []):
        w = b.weight_lbs or Decimal("0")
        raw_rate = b.rate_per_lb or Decimal("0")
        m_rate = max(Decimal("0"), raw_rate - Decimal("1"))
        extra = b.extra_fee or Decimal("0")
        tot = (w * m_rate) + extra
        bills_data.append({
            "stt": idx + 1,
            "consignee_name": b.consignee_name,
            "date": f"{b.bill_date.month}/{b.bill_date.day}/{b.bill_date.year}" if b.bill_date else "",
            "weight_lbs": float(w),
            "rate_per_lb": float(m_rate),
            "extra_fee": float(extra),
            "total_amount": float(tot),
        })

    date_str = f"{statement.statement_date.month}/{statement.statement_date.day}/{statement.statement_date.year}" if statement.statement_date else ""
    buf = generate_market_statement_pdf(
        statement_title=f"Bill Vận Chuyển — {statement.statement_code}",
        date_title=date_str,
        market_name=statement.market_name or "World Market",
        bills=bills_data,
        total_amount=float(statement.total_amount or 0),
        total_weight=float(statement.total_weight_lbs or 0),
        total_extra=float(statement.total_extra_fee or 0),
    )
    filename = f"Bill_{statement.statement_code}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ==================== SAVED STATEMENTS / BILLS CŨ ====================
@router.get("/statements", response_model=List[CargoStatementRead])
def list_statements(
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    """Lấy danh sách các bảng kê bill cũ đã xuất cho Chợ kèm trạng thái thanh toán."""
    statements = db.scalars(
        select(CargoStatement).order_by(desc(CargoStatement.statement_date)).offset(skip).limit(limit)
    ).all()
    return statements


@router.post("/statements", response_model=CargoStatementRead)
def save_statement(
    payload: CargoStatementCreate,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    """Lưu chốt một bảng kê mới cho Chợ và gán các bills vào bảng kê này."""
    if not payload.bill_ids:
        raise HTTPException(status_code=400, detail="Cần chọn ít nhất 1 bill để lập bảng kê")

    bills = db.scalars(select(CargoBill).where(and_(CargoBill.id.in_(payload.bill_ids), CargoBill.billing_type == "MARKET"))).all()
    if not bills:
        raise HTTPException(status_code=400, detail="Không tìm thấy các bills được chọn")

    total_weight = Decimal("0")
    total_extra = Decimal("0")
    total_amount = Decimal("0")

    for b in bills:
        w = b.weight_lbs or Decimal("0")
        market_rate = max(Decimal("0"), (b.rate_per_lb or Decimal("0")) - Decimal("1"))
        extra = b.extra_fee or Decimal("0")
        line_total = (w * market_rate) + extra

        total_weight += w
        total_extra += extra
        total_amount += line_total

    stmt_date = payload.statement_date or datetime.now(timezone.utc)
    code = f"BILL-{stmt_date.strftime('%Y%m%d')}-{uuid.uuid4().hex[:4].upper()}"

    statement = CargoStatement(
        statement_code=code,
        statement_date=stmt_date,
        market_name=payload.market_name or "World Market",
        total_weight_lbs=total_weight,
        total_extra_fee=total_extra,
        total_amount=total_amount,
        is_paid=False,
        notes=payload.notes,
    )
    db.add(statement)
    db.commit()
    db.refresh(statement)

    # Gán các bill vào statement này
    for b in bills:
        b.statement_id = statement.id
    db.commit()
    db.refresh(statement)

    return statement


@router.get("/statements/{statement_id}", response_model=CargoStatementRead)
def get_statement_detail(
    statement_id: int,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    statement = db.scalar(
        select(CargoStatement)
        .options(joinedload(CargoStatement.bills))
        .where(CargoStatement.id == statement_id)
    )
    if not statement:
        raise HTTPException(status_code=404, detail="Không tìm thấy bảng kê")
    return statement


@router.put("/statements/{statement_id}/toggle-paid")
def toggle_statement_paid(
    statement_id: int,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    statement = db.get(CargoStatement, statement_id)
    if not statement:
        raise HTTPException(status_code=404, detail="Không tìm thấy bảng kê")

    statement.is_paid = not statement.is_paid
    if statement.is_paid:
        statement.paid_at = datetime.now(timezone.utc)
        # Also mark all bills in this statement as paid
        for b in statement.bills:
            b.is_paid = True
    else:
        statement.paid_at = None
        for b in statement.bills:
            b.is_paid = False

    db.commit()
    db.refresh(statement)
    return {
        "status": "ok",
        "statement_id": statement.id,
        "is_paid": statement.is_paid,
        "paid_at": statement.paid_at.isoformat() if statement.paid_at else None,
    }


@router.delete("/statements/{statement_id}")
def delete_statement(
    statement_id: int,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    statement = db.get(CargoStatement, statement_id)
    if not statement:
        raise HTTPException(status_code=404, detail="Không tìm thấy bảng kê")

    # Unlink bills so they return to unbilled state
    for b in statement.bills:
        b.statement_id = None

    db.delete(statement)
    db.commit()
    return {"status": "ok", "message": "Đã xoá bảng kê, các bills đã trở lại trạng thái chưa chốt"}



# ==============================================================================
# PIRATE SHIP SHIPPING LABELS EXPORT
# ==============================================================================

import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

def format_box_recipients(names: list[str], max_len: int = 25) -> str:
    """
    Formats recipient names into a compact string:
    - 1 recipient: Full name if <= max_len, or last 2 words / given name.
    - Multiple recipients: 'A+B+C+D' (compact, no spaces around +).
    """
    if not names:
        return ""
    cleaned = [n.strip() for n in names if n.strip()]
    if not cleaned:
        return ""

    if len(cleaned) == 1:
        name = cleaned[0].upper()
        if len(name) <= max_len:
            return name
        parts = name.split()
        if len(parts) >= 2 and len(parts[-2] + " " + parts[-1]) <= max_len:
            return f"{parts[-2]} {parts[-1]}"
        return parts[-1]

    # Multiple recipients: A+B+C+D (compact without spaces)
    short_names = []
    for n in cleaned:
        parts = n.split()
        short_names.append(parts[-1].upper() if parts else n.upper())

    joined = "+".join(short_names)
    if len(joined) <= max_len:
        return joined

    # If still too long, pack as many as possible
    res = []
    curr_len = 0
    for i, sn in enumerate(short_names):
        sep = "+" if res else ""
        suffix = f"(+{len(short_names) - i})"
        if curr_len + len(sep) + len(sn) + len(suffix) <= max_len:
            res.append(sn)
            curr_len += len(sep) + len(sn)
        else:
            if res:
                return "+".join(res) + f"(+{len(short_names) - len(res)})"
            else:
                return sn[:max_len]
    return "+".join(res)


def extract_box_recipients(box: CargoBox) -> list[str]:
    """Extract recipient names from attached bills or box.label_text."""
    names = []
    seen = set()

    # 1. From attached bills
    if getattr(box, "bills", None):
        for bill in box.bills:
            cname = (bill.consignee_name or "").strip()
            if cname and cname.upper() not in seen:
                seen.add(cname.upper())
                names.append(cname)

    # 2. If no bills, parse box.label_text
    if not names and getattr(box, "label_text", None):
        text = box.label_text.strip()
        if text:
            parts = re.split(r"[,+/&\n]|(?:\s+và\s+)", text, flags=re.IGNORECASE)
            for p in parts:
                p = p.strip()
                if p and p.upper() not in seen:
                    seen.add(p.upper())
                    names.append(p)

    # Clean up: if single name is just generic 'THÙNG...' or matches box_code, ignore
    if len(names) == 1:
        val = names[0].strip().upper()
        if val.startswith("THÙNG") or val.startswith("THUNG") or (box.box_code and val == box.box_code.strip().upper()):
            names = []

    return names


class ExportShippingLabelsRequest(BaseModel):
    batch_id: Optional[int] = None
    box_count: Optional[int] = 5
    weight_lbs: Optional[float] = 8.0
    prefix: Optional[str] = "TC"
    format: Optional[str] = "xlsx"

@router.post("/export-shipping-labels")
def export_shipping_labels(
    payload: ExportShippingLabelsRequest,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    date_str = datetime.now().strftime("%Y%m%d")
    prefix = (payload.prefix or "TC").strip().upper()

    recipient_full_name = "KING CARGO - THAI MISSOURI"
    address_1 = "10471 Bolsa Ave"
    address_2 = ""
    city = "Westminster"
    state = "CA"
    zip_code = "92683-6757"
    country = "US"

    rows = []
    batch = None
    if payload.batch_id:
        batch = db.scalar(
            select(CargoBatch)
            .options(
                joinedload(CargoBatch.boxes).joinedload(CargoBox.bills),
                joinedload(CargoBatch.bills),
            )
            .where(CargoBatch.id == payload.batch_id)
        )

    if batch and batch.boxes:
        sorted_boxes = sorted(batch.boxes, key=lambda b: b.id)
        total_boxes = len(sorted_boxes)
        for idx, bx in enumerate(sorted_boxes, start=1):
            order_id = f"{batch.batch_code}-{bx.box_code}"

            # Format box title: T1 - A+B+C+D
            prefix_str = f"T{idx}"
            avail_len = max(10, 30 - len(prefix_str) - 3)
            if bx.is_hub and bx.hub_consignee_name:
                hub_name = bx.hub_consignee_name.strip().upper()
                recipients_str = hub_name[:avail_len]
            else:
                box_names = extract_box_recipients(bx)
                recipients_str = format_box_recipients(box_names, max_len=avail_len) if box_names else ""

            if recipients_str:
                box_title = f"{prefix_str} - {recipients_str}"[:30]
            else:
                box_title = f"{prefix_str}" if total_boxes == 1 else f"{prefix_str}/{total_boxes}"

            company_field = box_title
            stamp_1 = box_title
            # Khoá cứng khối lượng chuẩn 8.0 lbs (hoặc payload.weight_lbs) cho nhãn Pirate Ship
            w = float(payload.weight_lbs) if payload.weight_lbs and payload.weight_lbs > 0 else 8.0

            rows.append([
                order_id,
                recipient_full_name,
                company_field,
                address_1,
                address_2,
                city,
                state,
                zip_code,
                country,
                w,
                stamp_1,
            ])
        filename_base = f"pirateship-{batch.batch_code}-{total_boxes}boxes"
    else:
        box_count = max(1, min(payload.box_count or 5, 100))
        default_weight = payload.weight_lbs if payload.weight_lbs and payload.weight_lbs > 0 else 8.0
        for i in range(1, box_count + 1):
            order_id = f"{prefix}-{date_str}-{i:02d}"
            box_title = f"T{i}" if box_count == 1 else f"T{i}/{box_count}"
            company_field = box_title
            stamp_1 = box_title

            rows.append([
                order_id,
                recipient_full_name,
                company_field,
                address_1,
                address_2,
                city,
                state,
                zip_code,
                country,
                default_weight,
                stamp_1,
            ])
        filename_base = f"pirateship-trustcargo-{date_str}-{box_count}boxes"

    headers = [
        "Order ID",
        "Full Name",
        "Company",
        "Address Line 1",
        "Address Line 2",
        "City",
        "State",
        "Zip",
        "Country",
        "Weight (lbs)",
        "Rubber Stamp 1",
    ]

    if payload.format and payload.format.lower() == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        output.seek(0)
        
        return StreamingResponse(
            io.BytesIO(output.getvalue().encode("utf-8")),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename_base}.csv"}
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PirateShip Labels"

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for row_idx, r in enumerate(rows, start=2):
        ws.append(r)
        for col_idx in range(1, len(r) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            if col_idx in [1, 7, 8, 9, 10, 11, 12]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename_base}.xlsx"}
    )


@router.get("/export-shipping-labels")
def export_shipping_labels_get(
    request: Request,
    batch_id: Optional[int] = None,
    box_count: Optional[int] = 5,
    weight_lbs: float = 8.0,
    prefix: str = "TC",
    format: str = "xlsx",
    token: Optional[str] = None,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    return export_shipping_labels(
        ExportShippingLabelsRequest(
            batch_id=batch_id,
            box_count=box_count,
            weight_lbs=weight_lbs,
            prefix=prefix,
            format=format
        ),
        db=db,
        user=user
    )


# ==================== CALI RECONCILIATION & MASTER MANIFEST ====================
@router.post("/reconcile-cali", response_model=CaliReconcileResponse)
def reconcile_cali(
    payload: CaliReconcileRequest,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    query = select(CargoBill).options(joinedload(CargoBill.box))
    if payload.batch_ids:
        query = query.where(CargoBill.batch_id.in_(payload.batch_ids))
    elif payload.batch_id:
        query = query.where(CargoBill.batch_id == payload.batch_id)
    else:
        # Mặc định gom tất cả các bill đang chờ đối soát (PENDING) qua nhiều đợt ship (4-5 đợt gần nhất)
        query = query.where(CargoBill.reconcile_status == "PENDING")

    warehouse_bills = db.scalars(query).all()
    warehouse_map = {}
    for b in warehouse_bills:
        norm = normalize_name(b.consignee_name)
        warehouse_map.setdefault(norm, []).append(b)

    results = []
    matched_bill_ids = set()
    total_cali_amt = Decimal("0")
    total_wh_amt = Decimal("0")
    matched_count = 0
    weight_diff_count = 0
    missing_in_wh_count = 0

    boxes_query = select(CargoBox).options(joinedload(CargoBox.bills))
    if payload.batch_ids:
        boxes_query = boxes_query.where(CargoBox.batch_id.in_(payload.batch_ids))
    elif payload.batch_id:
        boxes_query = boxes_query.where(CargoBox.batch_id == payload.batch_id)
    else:
        pending_batch_ids = list({b.batch_id for b in warehouse_bills})
        if pending_batch_ids:
            boxes_query = boxes_query.where(CargoBox.batch_id.in_(pending_batch_ids))
        else:
            boxes_query = boxes_query.limit(0)
    all_boxes = db.scalars(boxes_query).unique().all()

    for item in payload.items:
        total_cali_amt += item.amount
        c_norm = normalize_name(item.consignee_name)
        matched_bill = None

        # Check if item matches a HUB Box (consolidated to VN agent)
        hub_matched_box = None
        for bx in all_boxes:
            if (bx.is_hub or bx.hub_consignee_name) and bx.bills:
                h_name = bx.hub_consignee_name or bx.label_text or ""
                norm_h = normalize_name(h_name)
                if norm_h and (c_norm == norm_h or c_norm in norm_h or norm_h in c_norm):
                    hub_matched_box = bx
                    break

        if hub_matched_box:
            total_sub_weight = sum((b.weight_lbs for b in hub_matched_box.bills), Decimal("0"))
            total_sub_amt = sum((b.total_amount for b in hub_matched_box.bills), Decimal("0"))
            diff = item.weight_lbs - total_sub_weight
            diff_abs = abs(diff)
            status_str = "WEIGHT_DIFF" if diff_abs > Decimal("1.0") else "MATCHED"
            if status_str == "WEIGHT_DIFF":
                weight_diff_count += 1
            else:
                matched_count += 1

            total_wh_amt += total_sub_amt
            sub_recs = []
            for b in hub_matched_box.bills:
                matched_bill_ids.add(b.id)
                if payload.auto_apply:
                    b.cali_tracking_code = item.tracking_code
                    b.reconcile_status = status_str
                b_icon = "🏪" if b.billing_type == "MARKET" else ("👤" if b.billing_type == "DIRECT" else "📦")
                sub_recs.append(f"{b_icon} {b.consignee_name} ({b.weight_lbs} lbs - {b.billing_type})")

            hub_title = hub_matched_box.hub_consignee_name or hub_matched_box.label_text
            results.append(CaliReconcileResultItem(
                cali_tracking=item.tracking_code,
                cali_name=item.consignee_name,
                cali_weight=item.weight_lbs,
                cali_amount=item.amount,
                matched_bill_id=hub_matched_box.bills[0].id if hub_matched_box.bills else None,
                warehouse_name=f"[ĐẦU CẦU VN] {hub_title} ({len(hub_matched_box.bills)} khách con)",
                warehouse_weight=total_sub_weight,
                warehouse_box=hub_matched_box.box_code,
                billing_type="HUB_CONSOLIDATED",
                weight_diff=diff,
                is_hub_box=True,
                sub_recipients=sub_recs,
                status=status_str
            ))
            continue

        if c_norm in warehouse_map and warehouse_map[c_norm]:
            matched_bill = warehouse_map[c_norm].pop(0)
        else:
            c_parts = [p for p in c_norm.split() if len(p) >= 2]
            for key, b_list in list(warehouse_map.items()):
                if b_list and (c_norm in key or key in c_norm or any(p in key for p in c_parts)):
                    matched_bill = b_list.pop(0)
                    break

        if matched_bill:
            matched_bill_ids.add(matched_bill.id)
            total_wh_amt += matched_bill.total_amount
            diff = item.weight_lbs - matched_bill.weight_lbs
            diff_abs = abs(diff)
            status_str = "WEIGHT_DIFF" if diff_abs > Decimal("0.5") else "MATCHED"
            if status_str == "WEIGHT_DIFF":
                weight_diff_count += 1
            else:
                matched_count += 1

            if payload.auto_apply:
                matched_bill.cali_tracking_code = item.tracking_code
                matched_bill.cali_weight_lbs = item.weight_lbs
                matched_bill.cali_amount = item.amount
                matched_bill.reconcile_status = status_str

            box_code = matched_bill.box.box_code if matched_bill.box else "CHƯA GHÉP"

            results.append(CaliReconcileResultItem(
                cali_tracking=item.tracking_code,
                cali_name=item.consignee_name,
                cali_weight=item.weight_lbs,
                cali_amount=item.amount,
                matched_bill_id=matched_bill.id,
                warehouse_name=matched_bill.consignee_name,
                warehouse_weight=matched_bill.weight_lbs,
                warehouse_box=box_code,
                billing_type=matched_bill.billing_type,
                weight_diff=diff,
                status=status_str
            ))
        else:
            missing_in_wh_count += 1
            results.append(CaliReconcileResultItem(
                cali_tracking=item.tracking_code,
                cali_name=item.consignee_name,
                cali_weight=item.weight_lbs,
                cali_amount=item.amount,
                matched_bill_id=None,
                warehouse_name=None,
                warehouse_weight=None,
                warehouse_box=None,
                billing_type=None,
                weight_diff=None,
                status="MISSING_IN_WAREHOUSE"
            ))

    if payload.auto_apply:
        db.commit()

    unmatched_wh = [CargoBillRead.model_validate(b) for b in warehouse_bills if b.id not in matched_bill_ids]

    return CaliReconcileResponse(
        total_cali_rows=len(payload.items),
        matched_rows=matched_count,
        weight_diff_rows=weight_diff_count,
        missing_in_warehouse_rows=missing_in_wh_count,
        warehouse_bills_unmatched=len(unmatched_wh),
        total_cali_amount=total_cali_amt,
        total_warehouse_amount=total_wh_amt,
        details=results,
        unmatched_warehouse_bills=unmatched_wh
    )


@router.get("/manifest")
def get_master_manifest(
    batch_id: Optional[int] = None,
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    query = select(CargoBill).options(joinedload(CargoBill.box), joinedload(CargoBill.batch))
    if batch_id:
        query = query.where(CargoBill.batch_id == batch_id)
    else:
        latest = db.scalar(select(CargoBatch).order_by(desc(CargoBatch.id)).limit(1))
        if latest:
            query = query.where(CargoBill.batch_id == latest.id)

    bills = db.scalars(query.order_by(CargoBill.id)).all()
    rows = []
    for idx, b in enumerate(bills, start=1):
        box_str = b.box.box_code if b.box else "CHƯA GHÉP"
        box_type = b.box.box_type.value if b.box else "SINGLE"
        rows.append({
            "stt": idx,
            "bill_id": b.id,
            "box_code": box_str,
            "box_type": box_type,
            "consignee_name": b.consignee_name,
            "consignee_phone": b.consignee_phone,
            "consignee_address": b.consignee_address,
            "weight_lbs": float(b.weight_lbs),
            "billing_type": b.billing_type,
            "shipper_name": b.shipper_name,
            "shipper_phone": b.shipper_phone,
            "packed_date": b.packed_date.isoformat() if b.packed_date else None,
            "bill_date": b.bill_date.isoformat() if b.bill_date else None,
            "rate_per_lb": float(b.rate_per_lb),
            "extra_fee": float(b.extra_fee),
            "total_amount": float(b.total_amount),
            "cali_tracking_code": b.cali_tracking_code,
            "cali_weight_lbs": float(b.cali_weight_lbs) if b.cali_weight_lbs else None,
            "cali_amount": float(b.cali_amount) if b.cali_amount else None,
            "reconcile_status": b.reconcile_status,
            "bill_image_url": b.bill_image_url,
        })
    return {
        "total_lines": len(rows),
        "total_weight_lbs": sum(r["weight_lbs"] for r in rows),
        "total_amount": sum(r["total_amount"] for r in rows),
        "lines": rows
    }


# ==================== DASHBOARD SUMMARY ENDPOINT ====================
@router.get("/dashboard-summary", response_model=TrustCargoDashboardSummary)
def get_dashboard_summary(
    db: Session = Depends(get_db),
    user: Optional[User] = Depends(get_current_user_or_bot),
):
    batches = db.scalars(
        select(CargoBatch)
        .options(
            joinedload(CargoBatch.boxes).joinedload(CargoBox.bills),
            joinedload(CargoBatch.bills),
        )
        .order_by(CargoBatch.ship_date.desc(), CargoBatch.id.desc())
    ).unique().all()

    all_bills = db.scalars(
        select(CargoBill)
        .options(joinedload(CargoBill.batch), joinedload(CargoBill.box))
        .order_by(CargoBill.created_at.desc())
    ).all()

    market_map = {}
    direct_map = {}
    self_list = []

    kpi = DashboardKpis()
    kpi.total_batches = len(batches)

    total_boxes_set = set()
    hub_boxes_set = set()

    for bill in all_bills:
        kpi.total_bills += 1
        w = bill.weight_lbs or Decimal("0")
        tot = bill.total_amount or Decimal("0")
        b_type = (bill.billing_type or "MARKET").upper()
        
        kpi.total_weight_lbs += w
        
        if bill.box_id:
            total_boxes_set.add(bill.box_id)
            if bill.box and getattr(bill.box, "is_hub", False):
                hub_boxes_set.add(bill.box_id)

        # Categorize
        if b_type == "MARKET":
            raw_rate = bill.rate_per_lb or Decimal("0")
            m_rate = max(Decimal("0"), raw_rate - Decimal("1"))
            extra = bill.extra_fee or Decimal("0")
            line_market_total = (w * m_rate) + extra
            tot = line_market_total

            kpi.market_revenue += tot
            kpi.total_revenue += tot
            unpaid = Decimal("0") if bill.is_paid else tot
            kpi.market_unpaid_balance += unpaid
            kpi.total_unpaid_balance += unpaid

            m_name = (bill.market_name or "World Market").strip()
            if m_name not in market_map:
                market_map[m_name] = {
                    "market_name": m_name,
                    "bill_count": 0,
                    "total_weight_lbs": Decimal("0"),
                    "total_amount": Decimal("0"),
                    "paid_amount": Decimal("0"),
                    "balance": Decimal("0"),
                    "bills": [],
                }
            m_data = market_map[m_name]
            m_data["bill_count"] += 1
            m_data["total_weight_lbs"] += w
            m_data["total_amount"] += tot
            if bill.is_paid:
                m_data["paid_amount"] += tot
            else:
                m_data["balance"] += tot
            b_read = CargoBillRead.model_validate(bill)
            b_read.total_amount = tot
            m_data['bills'].append(b_read)

        elif b_type == "DIRECT":
            kpi.direct_revenue += tot
            kpi.total_revenue += tot
            unpaid = Decimal("0") if bill.is_paid else tot
            kpi.direct_unpaid_balance += unpaid
            kpi.total_unpaid_balance += unpaid

            s_name = (bill.shipper_name or "Khách ngoài").strip()
            if s_name not in direct_map:
                direct_map[s_name] = {
                    "shipper_name": s_name,
                    "shipper_phone": bill.shipper_phone or "",
                    "bill_count": 0,
                    "total_weight_lbs": Decimal("0"),
                    "total_amount": Decimal("0"),
                    "paid_amount": Decimal("0"),
                    "balance": Decimal("0"),
                    "bills": [],
                }
            d_data = direct_map[s_name]
            d_data["bill_count"] += 1
            d_data["total_weight_lbs"] += w
            d_data["total_amount"] += tot
            if not d_data["shipper_phone"] and bill.shipper_phone:
                d_data["shipper_phone"] = bill.shipper_phone
            if bill.is_paid:
                d_data["paid_amount"] += tot
            else:
                d_data["balance"] += tot
            d_data["bills"].append(CargoBillRead.model_validate(bill))

        elif b_type == "SELF":
            kpi.self_packages_count += 1
            kpi.self_total_weight_lbs += w
            self_list.append(
                DashboardSelfItem(
                    bill_id=bill.id,
                    batch_id=bill.batch_id,
                    batch_code=bill.batch.batch_code if bill.batch else f"ID-{bill.batch_id}",
                    box_id=bill.box_id,
                    box_code=bill.box.box_code if bill.box else None,
                    packed_date=bill.packed_date or bill.created_at,
                    consignee_name=bill.consignee_name or "",
                    consignee_phone=bill.consignee_phone or "",
                    consignee_address=bill.consignee_address or "",
                    commodities=bill.commodities or "",
                    weight_lbs=w,
                    cali_tracking_code=bill.cali_tracking_code,
                    cali_weight_lbs=bill.cali_weight_lbs,
                    cali_amount=bill.cali_amount,
                    notes=bill.notes,
                )
            )

    kpi.total_boxes = len(total_boxes_set)
    kpi.hub_boxes_count = len(hub_boxes_set)

    # 2. Build shipment logs grouped by batch
    shipment_logs = []
    for b in batches:
        boxes_log = []
        b_weight = Decimal("0")
        b_bill_count = 0
        
        for bx in sorted(b.boxes, key=lambda x: x.id):
            bx_bills = [CargoBillRead.model_validate(bi) for bi in sorted(bx.bills, key=lambda x: x.id)]
            if bx.weight_lbs:
                bx_weight = bx.weight_lbs
            else:
                bx_weight = sum((bi.weight_lbs or Decimal("0") for bi in bx_bills), Decimal("0"))
            b_weight += bx_weight
            b_bill_count += len(bx_bills)
            
            boxes_log.append(
                DashboardShipmentLogBox(
                    id=bx.id,
                    box_code=bx.box_code or f"T{bx.id}",
                    box_type=str(bx.box_type.value) if hasattr(bx.box_type, "value") else str(bx.box_type),
                    is_hub=bool(getattr(bx, "is_hub", False)),
                    hub_consignee_name=getattr(bx, "hub_consignee_name", None),
                    weight_lbs=bx.weight_lbs,
                    label_text=bx.label_text or "",
                    bills=bx_bills,
                )
            )

        shipment_logs.append(
            DashboardShipmentLogBatch(
                id=b.id,
                batch_code=b.batch_code,
                ship_date=b.ship_date,
                market_name=b.market_name,
                status=str(b.status.value) if hasattr(b.status, "value") else str(b.status),
                box_count=len(boxes_log),
                bill_count=b_bill_count,
                total_weight_lbs=b_weight,
                boxes=boxes_log,
            )
        )

    return TrustCargoDashboardSummary(
        kpis=kpi,
        shipment_logs=shipment_logs,
        market_billing=[DashboardMarketItem(**v) for v in market_map.values()],
        direct_billing=[DashboardDirectCustomer(**v) for v in direct_map.values()],
        self_billing=self_list,
    )
